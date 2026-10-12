from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import unittest


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

import clean_academic_document as cleaner


class StructureCleaningTests(unittest.TestCase):
    def test_keeps_abstract_body_and_references_only(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "一种中文文本检测方法", kind_hint="heading"),
            cleaner.SourceBlock(1, "张三 李四 某某大学"),
            cleaner.SourceBlock(2, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(3, "本文提出一种中文文本检测方法，并在公开数据集上进行了验证。"),
            cleaner.SourceBlock(4, "关键词：人工智能；文本检测"),
            cleaner.SourceBlock(5, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(6, "近年来生成模型快速发展，学术文本的来源识别受到关注。"),
            cleaner.SourceBlock(7, "图1 系统总体结构", kind_hint="caption"),
            cleaner.SourceBlock(8, "模型 | AUROC\nA | 0.90", kind_hint="table"),
            cleaner.SourceBlock(9, "- 第一项贡献", kind_hint="list", is_list=True),
            cleaner.SourceBlock(10, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(11, "[1] 张三. 中文文本检测研究[J]. 计算机学报, 2020."),
            cleaner.SourceBlock(12, "附录A", kind_hint="heading"),
            cleaner.SourceBlock(13, "这里是附录中的补充实验数据和说明。"),
        ]

        result, warnings = cleaner.clean_blocks(blocks)
        included = [item for item in result if item.included]

        self.assertEqual(
            [item.section_type for item in included],
            ["abstract", "body", "references"],
        )
        self.assertIn("本文提出一种中文文本检测方法", included[0].text)
        self.assertIn("近年来生成模型快速发展", included[1].text)
        self.assertTrue(included[2].text.startswith("[1]"))
        self.assertFalse(any("附录中的" in item.text for item in included))
        self.assertEqual(warnings, [])

    def test_inline_abstract_and_clean_offsets(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "题目", kind_hint="heading"),
            cleaner.SourceBlock(1, "摘要：这是摘要的完整正文，用于验证行内摘要识别。"),
            cleaner.SourceBlock(2, "一、引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "这是正文的连续自然语言段落，用于验证字符位置。"),
            cleaner.SourceBlock(4, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(5, "[1] Author. Paper title. Journal, 2020."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        clean_text = cleaner._assign_clean_offsets(result)
        included = [item for item in result if item.included]

        self.assertEqual(clean_text[included[0].clean_start : included[0].clean_end], included[0].text)
        self.assertEqual(clean_text[included[-1].clean_start : included[-1].clean_end], included[-1].text)
        self.assertNotIn("摘要：", clean_text)
        self.assertNotIn("一、引言", clean_text)

    def test_reference_numbering_is_not_treated_as_list(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "该摘要包含足够长度的自然语言文本以通过清洗。"),
            cleaner.SourceBlock(2, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "该正文包含足够长度的自然语言文本以通过清洗。"),
            cleaner.SourceBlock(4, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(5, "[1] Test reference.", kind_hint="list", is_list=True),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        reference = next(item for item in result if item.text.startswith("[1]"))
        self.assertTrue(reference.included)
        self.assertEqual(reference.block_type, "reference")

    def test_numbered_reference_does_not_restart_body(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(2, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "这是应当保留的正文连续自然语言段落。"),
            cleaner.SourceBlock(4, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(5, "140. IEEE Standard for Example Data, 2021."),
            cleaner.SourceBlock(6, "附录A", kind_hint="heading"),
            cleaner.SourceBlock(7, "附录内容不应保留。"),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        reference = next(item for item in result if item.text.startswith("140."))
        appendix = next(item for item in result if item.text.startswith("附录内容"))
        self.assertTrue(reference.included)
        self.assertEqual(reference.section_type, "references")
        self.assertFalse(appendix.included)

    def test_unnumbered_references_are_split_by_terminal_year(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(2, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "这是应当保留的正文连续自然语言段落。"),
            cleaner.SourceBlock(4, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(
                5,
                "Author A. First paper. Journal, 2020. Author B. Second paper. Journal, 2021.",
            ),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        references = [item for item in result if item.included and item.section_type == "references"]
        self.assertEqual(len(references), 2)
        self.assertTrue(references[0].text.startswith("Author A"))
        self.assertTrue(references[1].text.startswith("Author B"))

    def test_lettered_appendix_after_references_is_excluded(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(2, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "这是应当保留的正文连续自然语言段落。"),
            cleaner.SourceBlock(4, "REFERENCES", kind_hint="heading"),
            cleaner.SourceBlock(5, "Author A. First paper. Journal, 2020."),
            cleaner.SourceBlock(6, "A EXPERIMENTAL SETTINGS", kind_hint="heading"),
            cleaner.SourceBlock(7, "1 ADDITIONAL RESULTS", kind_hint="heading"),
            cleaner.SourceBlock(8, "Appendix prose should not be included in detection."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        appendix = next(item for item in result if item.text.startswith("Appendix prose"))
        self.assertFalse(appendix.included)
        self.assertEqual(appendix.exclude_reason, "appendix")

    def test_heading_styled_title_does_not_start_body(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "论文题名", kind_hint="heading"),
            cleaner.SourceBlock(1, "张三李四王五某某大学计算机学院"),
            cleaner.SourceBlock(2, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(3, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(4, "Introduction", kind_hint="heading"),
            cleaner.SourceBlock(5, "This body paragraph contains enough words for the cleaner."),
            cleaner.SourceBlock(6, "References", kind_hint="heading"),
            cleaner.SourceBlock(7, "[1] Author. Reference title. Journal, 2020."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        author = next(item for item in result if item.text.startswith("张三李四"))
        self.assertFalse(author.included)
        self.assertEqual(author.exclude_reason, "front_matter")

    def test_equation_and_metadata_are_excluded(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "本文对数据进行了分析，并报告了主要实验结果。"),
            cleaner.SourceBlock(2, "1 方法", kind_hint="heading"),
            cleaner.SourceBlock(3, r"L=\\sum_{i=1}^{n} x_i^2"),
            cleaner.SourceBlock(4, "基金项目：某某基金资助"),
            cleaner.SourceBlock(5, "本节介绍模型训练过程以及采用的评价指标。"),
            cleaner.SourceBlock(6, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(7, "[1] Test. Reference, 2021."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        by_text = {item.text: item for item in result}
        self.assertEqual(by_text[r"L=\\sum_{i=1}^{n} x_i^2"].exclude_reason, "equation_or_code")
        self.assertEqual(by_text["基金项目：某某基金资助"].exclude_reason, "publication_metadata")
        self.assertTrue(by_text["本节介绍模型训练过程以及采用的评价指标。"].included)

    def test_inline_markdown_math_in_prose_is_kept_for_rule_detection(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(2, "1 方法", kind_hint="heading"),
            cleaner.SourceBlock(
                3,
                "本文采用目标函数$L=N/D$计算损失，并对实验结果进行比较分析。",
            ),
            cleaner.SourceBlock(4, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(5, "[1] Test. Reference, 2021."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        target = next(item for item in result if "$L=N/D$" in item.text)
        self.assertTrue(target.included)
        self.assertEqual(target.block_type, "paragraph")

    def test_caption_is_excluded_but_figure_explanation_is_kept(self) -> None:
        blocks = [
            cleaner.SourceBlock(0, "摘要", kind_hint="heading"),
            cleaner.SourceBlock(1, "这是应当保留的摘要连续自然语言正文。"),
            cleaner.SourceBlock(2, "1 引言", kind_hint="heading"),
            cleaner.SourceBlock(3, "图 4-3 各质量维度内指标的相关系数热力图"),
            cleaner.SourceBlock(
                4,
                "图 4-3 展示了各维度内部指标的相关关系，说明部分指标存在信息重叠。",
            ),
            cleaner.SourceBlock(5, "参考文献", kind_hint="heading"),
            cleaner.SourceBlock(6, "[1] Author. Paper title. Journal, 2020."),
        ]

        result, _ = cleaner.clean_blocks(blocks)
        by_text = {item.text: item for item in result}
        self.assertFalse(by_text["图 4-3 各质量维度内指标的相关系数热力图"].included)
        self.assertTrue(by_text["图 4-3 展示了各维度内部指标的相关关系，说明部分指标存在信息重叠。"].included)


class OutputTests(unittest.TestCase):
    def test_clean_document_accepts_docx_and_excludes_table(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "paper.docx"
            document = Document()
            document.add_heading("论文题名", level=1)
            document.add_heading("摘要", level=1)
            document.add_paragraph("这是应当保留的摘要连续自然语言正文。")
            document.add_paragraph("关键词：生成式人工智能；文本检测")
            document.add_heading("1 引言", level=1)
            document.add_paragraph("这是应当保留的正文连续自然语言段落。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "模型"
            table.cell(0, 1).text = "AUROC"
            table.cell(1, 0).text = "A"
            table.cell(1, 1).text = "0.90"
            document.add_heading("参考文献", level=1)
            document.add_paragraph("[1] Author. Paper title. Journal, 2020.")
            document.save(source_path)

            result = cleaner.clean_document(source_path)

        self.assertEqual(result["source"]["format"], "docx")
        self.assertIn("应当保留的摘要", result["clean_text"])
        self.assertIn("应当保留的正文", result["clean_text"])
        self.assertIn("[1] Author", result["clean_text"])
        self.assertNotIn("AUROC", result["clean_text"])

    def test_clean_document_rejects_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "paper.pdf"
            source_path.write_bytes(b"%PDF-1.4")
            with self.assertRaisesRegex(cleaner.CleaningError, "PDF 拆解由上游系统负责"):
                cleaner.clean_document(source_path)

    def test_inline_formula_is_marked_and_formula_only_paragraph_is_excluded(self) -> None:
        from docx import Document
        from docx.oxml import OxmlElement

        def append_math(paragraph, expression: str) -> None:
            math = OxmlElement("m:oMath")
            math_run = OxmlElement("m:r")
            math_text = OxmlElement("m:t")
            math_text.text = expression
            math_run.append(math_text)
            math.append(math_run)
            paragraph._p.append(math)

        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "formula.docx"
            document = Document()
            document.add_heading("摘要", level=1)
            document.add_paragraph("这是应当保留的摘要连续自然语言正文。")
            document.add_heading("1 引言", level=1)
            inline = document.add_paragraph("参数")
            append_math(inline, "theta")
            inline.add_run("越大时，模型的预测损失越低。")
            formula_only = document.add_paragraph()
            append_math(formula_only, "L=E+AN")
            document.add_heading("参考文献", level=1)
            document.add_paragraph("[1] Author. Paper title. Journal, 2020.")
            document.save(source_path)

            result = cleaner.clean_document(source_path)

        self.assertIn("参数[公式]越大时", result["clean_text"])
        inline_block = next(block for block in result["blocks"] if "参数[公式]" in block["text"])
        self.assertTrue(inline_block["metadata"]["contains_inline_formula"])
        self.assertEqual(inline_block["metadata"]["inline_formula_count"], 1)
        formula_block = next(block for block in result["blocks"] if block["text"] == "[公式]")
        self.assertFalse(formula_block["included"])
        self.assertEqual(formula_block["exclude_reason"], "equation_or_code")

    def test_write_result_creates_utf8_files(self) -> None:
        result = {
            "clean_text": "摘要正文\n\n正文内容",
            "stats": {
                "included_block_count": 2,
                "excluded_block_count": 0,
                "included_char_count": 10,
            },
            "warnings": [],
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            text_path = Path(tmpdir) / "paper.cleaned.txt"
            json_path = Path(tmpdir) / "paper.cleaned.json"
            cleaner.write_result(result, text_path, json_path)
            self.assertIn("摘要正文", text_path.read_text(encoding="utf-8"))
            loaded = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertEqual(loaded["clean_text"], result["clean_text"])


if __name__ == "__main__":
    unittest.main()
