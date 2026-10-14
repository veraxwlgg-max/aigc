#!/usr/bin/env python3
"""在论文副本中分散注入 R02 的 27 类格式信号，用于端到端回归测试。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any


class InjectionError(RuntimeError):
    pass


INJECTIONS = (
    {
        "signal": "markdown_front_matter",
        "mode": "append_lines",
        "break_count": 1,
        "text": "---\ntitle: 格式残留测试\n---",
    },
    {
        "signal": "markdown_heading",
        "mode": "append_lines",
        "break_count": 1,
        "text": "## 格式测试小标题",
    },
    {
        "signal": "markdown_setext_heading",
        "mode": "append_lines",
        "break_count": 1,
        "text": "格式测试标题\n---",
    },
    {
        "signal": "markdown_bold_italic",
        "mode": "append",
        "text": " ***格式测试粗斜体***",
    },
    {
        "signal": "markdown_bold",
        "mode": "append",
        "text": " **格式测试加粗**",
    },
    {
        "signal": "markdown_italic",
        "mode": "append",
        "text": " *格式测试斜体*",
    },
    {
        "signal": "markdown_strikethrough",
        "mode": "append",
        "text": " ~~格式测试删除线~~",
    },
    {
        "signal": "markdown_blockquote",
        "mode": "append_lines",
        "break_count": 1,
        "text": "> [!NOTE] 格式测试提示块",
    },
    {
        "signal": "markdown_task_list",
        "mode": "append_lines",
        "break_count": 1,
        "text": "- [x] 格式测试任务列表",
    },
    {
        "signal": "bullet_list",
        "mode": "append_lines",
        "break_count": 1,
        "text": "- 格式测试第一点",
    },
    {
        "signal": "numbered_list",
        "mode": "append_lines",
        "break_count": 1,
        "text": "1. 格式测试编号项",
    },
    {
        "signal": "markdown_image",
        "mode": "append",
        "text": " ![格式测试图片](figure-test.png)",
    },
    {
        "signal": "markdown_link",
        "mode": "append",
        "text": " [格式测试链接](https://example.org/test)",
    },
    {
        "signal": "markdown_reference_link",
        "mode": "append",
        "text": " [格式测试主页][test-project]",
    },
    {
        "signal": "markdown_link_definition",
        "mode": "append_lines",
        "break_count": 1,
        "text": "[test-project]: https://example.org/project",
    },
    {
        "signal": "markdown_autolink",
        "mode": "append",
        "text": " <https://example.org/test>",
    },
    {
        "signal": "markdown_fenced_code",
        "mode": "append_lines",
        "break_count": 1,
        "text": "```python\nprint('format-test')\n```",
    },
    {
        "signal": "markdown_inline_code",
        "mode": "append",
        "text": " `detect_format()`",
    },
    {
        "signal": "markdown_table",
        "mode": "append_lines",
        "break_count": 1,
        "text": "| 字段 | 数值 |\n| --- | :---: |\n| A | 1 |",
    },
    {
        "signal": "markdown_horizontal_rule",
        "mode": "append_lines",
        "break_count": 1,
        "text": "***",
    },
    {
        "signal": "markdown_footnote",
        "mode": "append",
        "text": " 格式测试脚注[^format-test]",
    },
    {
        "signal": "markdown_raw_html",
        "mode": "append",
        "text": " <details><summary>格式测试</summary>残留内容</details>",
    },
    {
        "signal": "markdown_math_delimiter",
        "mode": "append",
        "text": " 格式测试公式$L=N/D$",
    },
    {
        "signal": "markdown_escape_sequence",
        "mode": "append",
        "text": r" 格式测试\*转义星号\*",
    },
    {
        "signal": "frequent_blank_lines",
        "mode": "append_lines",
        "break_count": 3,
        "text": "格式测试空行之后继续说明。",
    },
    {
        "signal": "answer_heading",
        "mode": "prepend",
        "text": "以下是根据您的要求整理的内容：",
    },
    {
        "signal": "closing_suggestion",
        "mode": "append",
        "text": " 如需进一步说明，我可以继续为您补充。",
    },
)


REFERENCE_HEADINGS = {"参考文献", "主要参考文献", "references", "bibliography"}
EXCLUDED_STYLE_WORDS = {
    "heading",
    "title",
    "toc",
    "caption",
    "list",
    "标题",
    "目录",
    "题注",
    "列表",
}


def is_candidate_paragraph(paragraph: object) -> bool:
    """判断段落是否适合承载格式注入，不读取或保存其具体内容。"""

    text = paragraph.text.strip()
    if len(text) < 80:
        return False
    visible_characters = sum(char.isalnum() for char in text)
    if visible_characters < 40:
        return False
    style = getattr(paragraph, "style", None)
    style_name = str(getattr(style, "name", "")).lower()
    return not any(word in style_name for word in EXCLUDED_STYLE_WORDS)


def select_target_paragraphs(document: object, count: int) -> list[tuple[int, object]]:
    """从参考文献之前的正文候选段中均匀选取目标段落。"""

    candidates: list[tuple[int, object]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text.lower() in REFERENCE_HEADINGS:
            break
        if is_candidate_paragraph(paragraph):
            candidates.append((index, paragraph))

    if len(candidates) < count:
        raise InjectionError(
            f"可用正文段落不足：需要 {count} 个，实际找到 {len(candidates)} 个。"
        )
    if count == 1:
        return [candidates[len(candidates) // 2]]

    positions = [index * (len(candidates) - 1) // (count - 1) for index in range(count)]
    return [candidates[position] for position in positions]


def prepend_text(paragraph: object, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text + paragraph.runs[0].text
    else:
        paragraph.add_run(text)


def inject(source: Path, output: Path) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise InjectionError("需要 python-docx 才能生成测试副本。") from exc

    if not source.is_file():
        raise InjectionError(f"原始 Word 文件不存在：{source}")
    if source.resolve() == output.resolve():
        raise InjectionError("输出路径不得覆盖原始 Word 文件。")

    document = Document(source)
    targets = select_target_paragraphs(document, len(INJECTIONS))
    manifest: list[dict[str, Any]] = []
    for item, (paragraph_index, paragraph) in zip(INJECTIONS, targets, strict=True):
        if item["mode"] == "prepend":
            prepend_text(paragraph, item["text"])
        elif item["mode"] == "append":
            paragraph.add_run(item["text"])
        elif item["mode"] == "append_lines":
            run = paragraph.add_run()
            for _ in range(int(item["break_count"])):
                run.add_break()
            run.add_text(item["text"])
        else:
            raise InjectionError(f"未知注入模式：{item['mode']}")
        manifest.append(
            {
                "signal": item["signal"],
                "paragraph_index": paragraph_index,
                "injected_text": item["text"],
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    manifest_path = output.with_suffix(".injections.json")
    manifest_path.write_text(
        json.dumps(
            {
                "source_name": source.name,
                "output_name": output.name,
                "expected_signals": [item["signal"] for item in manifest],
                "injections": manifest,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="原始 DOCX")
    parser.add_argument("output", type=Path, help="另存的测试 DOCX")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        manifest = inject(args.source, args.output)
    except (InjectionError, OSError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2

    print(f"已生成：{args.output}")
    print(f"注入清单：{args.output.with_suffix('.injections.json')}")
    for item in manifest:
        print(f"- {item['signal']}: {item['injected_text']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
